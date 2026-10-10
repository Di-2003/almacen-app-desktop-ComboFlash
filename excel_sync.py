"""
Generación del Excel completo (optimizado con lxml + estilos pre-creados).

Hojas:
  - Por local (con columnas de moneda)
  - Movimientos (todas las operaciones del período)
  - Ventas_<Local> (solo ventas por local)
  - VentasGenerales (agregado por mes)

Extras:
  - generar_excel_salidas_hoy(local_id): todas las SALIDAS de hoy
  - generar_excel_diario(): resumen + una hoja por local con movs del día

Performance:
  - Estilos pre-creados (no se instancian miles de objetos).
  - Consultas agrupadas (una sola query por bloque).
  - `lxml` instalado → openpyxl lo usa automáticamente 3-5x más rápido.
"""
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db import get_conn, GENERAL_ID
import locales as loc
import inventario as inv


COLOR_HEADER_BG = "FF263238"
COLOR_HEADER_FG = "FFFFFFFF"
COLOR_TITULO = "FF37474F"
COLOR_TOTAL_BG = "FFECEFF1"
COLOR_MERMA = "FFC00000"

# ══════════════════════════════════════════════════════════
# ESTILOS PRE-CREADOS (evita instanciar miles de objetos)
# ══════════════════════════════════════════════════════════

BORDE_FINO = Border(
    left=Side(style="thin", color="FFB0BEC5"),
    right=Side(style="thin", color="FFB0BEC5"),
    top=Side(style="thin", color="FFB0BEC5"),
    bottom=Side(style="thin", color="FFB0BEC5"),
)

_FONT_HEADER = Font(bold=True, color=COLOR_HEADER_FG, size=10)
_FONT_TITULO = Font(bold=True, size=12, color=COLOR_TITULO)
_FONT_BOLD = Font(bold=True)
_FONT_BOLD_MERMA = Font(bold=True, color=COLOR_MERMA)
_FONT_BOLD_TITULO = Font(bold=True, color=COLOR_TITULO)
_FONT_ITALIC_GREY = Font(italic=True, color="FF888888")
_FONT_MERMA = Font(color=COLOR_MERMA)

_FILL_HEADER = PatternFill("solid", fgColor=COLOR_HEADER_BG)
_FILL_TOTAL = PatternFill("solid", fgColor=COLOR_TOTAL_BG)

_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center", indent=1)
_ALIGN_LEFT_PLAIN = Alignment(horizontal="left", vertical="center")

_NUM_2DEC = '#,##0.00'

NOMBRES_MES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
               "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre",
               "Diciembre"]


# ══════════════════════════════════════════════════════════
# ENTRY PRINCIPAL
# ══════════════════════════════════════════════════════════

def generar_excel_bytes() -> bytes:
    wb = Workbook()
    usados = set()
    primera = True

    locales_activos = loc.listar_locales(solo_activos=True)
    locales_todos = loc.listar_locales(solo_activos=False)

    # ── Cargar todo de una sola vez ──
    todos_productos = _cargar_productos_por_local(locales_activos)
    movs = _cargar_movimientos()
    ventas_por_local = _cargar_ventas_por_local(locales_todos)

    # ── Hojas por local ──
    for local in locales_activos:
        nombre = _sanitizar(local["nombre"], usados)
        if primera:
            ws = wb.active
            ws.title = nombre
            primera = False
        else:
            ws = wb.create_sheet(nombre)
        _hoja_productos_local(ws, local,
                              todos_productos.get(local["id"], []))

    # ── Movimientos ──
    _hoja_movimientos(wb.create_sheet(_sanitizar("Movimientos", usados)),
                      movs)

    # ── Ventas por local ──
    for local in locales_todos:
        nombre = _sanitizar(f"Ventas_{local['nombre']}", usados)
        _hoja_ventas_local(
            wb.create_sheet(nombre), local,
            ventas_por_local.get(local["id"], []))

    # ── Ventas generales ──
    _hoja_ventas_generales(
        wb.create_sheet(_sanitizar("VentasGenerales", usados)))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════
# CARGA DE DATOS (una sola pasada)
# ══════════════════════════════════════════════════════════

def _cargar_productos_por_local(locales) -> dict:
    if not locales:
        return {}
    ids = [l["id"] for l in locales]
    placeholders = ",".join("?" * len(ids))
    sql = (f"SELECT * FROM productos WHERE local_id IN ({placeholders}) "
           f"AND activo=1 "
           f"ORDER BY "
           f"CASE WHEN codigo IS NULL OR codigo='' THEN 1 ELSE 0 END, "
           f"codigo COLLATE NOCASE, nombre COLLATE NOCASE")
    result = {lid: [] for lid in ids}
    with get_conn() as conn:
        for r in conn.execute(sql, tuple(ids)).fetchall():
            d = dict(r)
            result[d["local_id"]].append(d)
    return result


def _cargar_movimientos() -> list[dict]:
    with get_conn() as conn:
        filas = conn.execute(
            "SELECT m.fecha, m.tipo, m.cantidad, m.motivo, m.rebaja, "
            "m.precio_unitario_momento, m.detalle, "
            "p.nombre AS producto, p.codigo AS codigo, "
            "l.nombre AS local FROM movimientos m "
            "JOIN productos p ON p.id=m.producto_id "
            "JOIN locales l ON l.id=m.local_id "
            "WHERE m.tipo IN ('ENTRADA','SALIDA','BAJA',"
            "                 'TRASPASO_SALIDA','TRASPASO_ENTRADA') "
            "ORDER BY m.fecha ASC, m.id ASC"
        ).fetchall()
    return [dict(r) for r in filas]


def _cargar_ventas_por_local(locales) -> dict:
    if not locales:
        return {}
    ids = [l["id"] for l in locales]
    placeholders = ",".join("?" * len(ids))
    sql = (f"SELECT m.local_id, m.fecha, m.cantidad, m.rebaja, "
           f"m.precio_unitario_momento, "
           f"p.nombre AS producto, p.codigo AS codigo "
           f"FROM movimientos m "
           f"JOIN productos p ON p.id=m.producto_id "
           f"WHERE m.local_id IN ({placeholders}) "
           f"AND m.tipo='SALIDA' "
           f"AND LOWER(TRIM(COALESCE(m.motivo,'')))='venta' "
           f"ORDER BY m.fecha ASC, m.id ASC")
    result = {lid: [] for lid in ids}
    with get_conn() as conn:
        for r in conn.execute(sql, tuple(ids)).fetchall():
            d = dict(r)
            result[d["local_id"]].append(d)
    return result


# ══════════════════════════════════════════════════════════
# HELPERS
# ══════════════════════════════════════════════════════════

def _r2(x):
    try:
        return round(float(x or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def _sanitizar(nombre, usados):
    prohibidos = set(':\\/?*[]')
    limpio = "".join(c for c in (nombre or "")
                     if c not in prohibidos).strip() or "Hoja"
    limpio = limpio[:31]
    base = limpio
    cont = 1
    while limpio.lower() in usados:
        sufijo = f"_{cont}"
        limpio = base[:31 - len(sufijo)] + sufijo
        cont += 1
    usados.add(limpio.lower())
    return limpio


def _solo_fecha(t):
    return (t or "").split(" ")[0]


def _encabezado(ws, fila, textos):
    for col, txt in enumerate(textos, 1):
        c = ws.cell(row=fila, column=col, value=txt)
        c.font = _FONT_HEADER
        c.fill = _FILL_HEADER
        c.alignment = _ALIGN_CENTER
        c.border = BORDE_FINO
    ws.row_dimensions[fila].height = 26


def _anchos(ws, anchos):
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _titulo_hoja(ws, col_fin, texto):
    ws["A1"] = texto
    ws["A1"].font = _FONT_TITULO
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=col_fin)
    ws.row_dimensions[1].height = 24


def _estilizar_fila(ws, fila, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=fila, column=col)
        c.border = BORDE_FINO
        c.fill = _FILL_TOTAL


def _cierre_dia(ws, fila, etiqueta, rebajas, venta, ncols,
                col_rebaja, col_venta, col_merma=None, merma=0.0):
    c = ws.cell(row=fila, column=1, value=etiqueta)
    c.font = _FONT_BOLD_TITULO
    c.alignment = _ALIGN_LEFT
    _estilizar_fila(ws, fila, ncols)
    fin_merge = max(1, col_rebaja - 1)
    if fin_merge >= 2:
        ws.merge_cells(start_row=fila, start_column=1,
                       end_row=fila, end_column=fin_merge)
    c = ws.cell(row=fila, column=col_rebaja, value=_r2(rebajas))
    c.font = _FONT_BOLD
    c.alignment = _ALIGN_RIGHT
    c.number_format = _NUM_2DEC
    c = ws.cell(row=fila, column=col_venta, value=_r2(venta))
    c.font = _FONT_BOLD
    c.alignment = _ALIGN_RIGHT
    c.number_format = _NUM_2DEC
    if col_merma is not None:
        c = ws.cell(row=fila, column=col_merma, value=_r2(merma))
        c.font = _FONT_BOLD_MERMA
        c.alignment = _ALIGN_RIGHT
        c.number_format = _NUM_2DEC


# ══════════════════════════════════════════════════════════
# HOJA POR LOCAL
# ══════════════════════════════════════════════════════════

def _hoja_productos_local(ws, local, productos):
    _titulo_hoja(
        ws, 11,
        f"{local['nombre']} — "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")

    FILA_HEADER = 3
    _encabezado(ws, FILA_HEADER, [
        "Código", "Producto", "Stock",
        "Moneda", "Costo orig.", "Costo CUP",
        "Venta orig.", "Venta CUP",
        "Costo Total", "Venta Total", "Margen",
    ])

    inv_total = venta_total = margen_total = 0.0
    fila = FILA_HEADER

    for p in productos:
        fila += 1
        stock = float(p["stock"] or 0)
        pc_cup = float(p.get("precio_costo", 0) or 0)
        pu_cup = float(p.get("precio_unitario", 0) or 0)
        pc_orig = float(p.get("precio_costo_orig", 0) or 0)
        pu_orig = float(p.get("precio_unitario_orig", 0) or 0)
        mc = (p.get("moneda_costo") or "CUP").upper()
        mv = (p.get("moneda_venta") or "CUP").upper()
        pc_tot = pc_cup * stock
        pv_tot = pu_cup * stock
        margen = pv_tot - pc_tot

        inv_total += pc_tot
        venta_total += pv_tot
        margen_total += margen

        moneda_txt = f"{mc}/{mv}" if mc != mv else mc
        pc_orig_mostrar = pc_orig if mc != "CUP" else pc_cup
        pu_orig_mostrar = pu_orig if mv != "CUP" else pu_cup

        valores = [
            p.get("codigo") or "",
            p["nombre"],
            stock,
            moneda_txt,
            _r2(pc_orig_mostrar),
            _r2(pc_cup),
            _r2(pu_orig_mostrar),
            _r2(pu_cup),
            _r2(pc_tot),
            _r2(pv_tot),
            _r2(margen),
        ]
        for col, val in enumerate(valores, 1):
            c = ws.cell(row=fila, column=col, value=val)
            c.border = BORDE_FINO
            if col in (1, 3, 4):
                c.alignment = _ALIGN_CENTER
            elif col >= 5:
                c.alignment = _ALIGN_RIGHT
                c.number_format = _NUM_2DEC
            else:
                c.alignment = _ALIGN_LEFT

    fila_tot = fila + 1
    c = ws.cell(row=fila_tot, column=1, value="TOTALES")
    c.font = _FONT_BOLD
    c.alignment = _ALIGN_LEFT
    for col, val in [(9, inv_total), (10, venta_total),
                     (11, margen_total)]:
        c = ws.cell(row=fila_tot, column=col, value=_r2(val))
        c.font = _FONT_BOLD
        c.alignment = _ALIGN_RIGHT
        c.number_format = _NUM_2DEC
    _estilizar_fila(ws, fila_tot, 11)

    ws["A2"] = (f"Invertido: ${inv_total:,.2f}   |   "
                f"Venta: ${venta_total:,.2f}   |   "
                f"Margen: ${margen_total:,.2f}")
    ws["A2"].font = Font(bold=True, size=10, color=COLOR_TITULO)
    ws.merge_cells("A2:K2")
    ws.row_dimensions[2].height = 20

    _anchos(ws, [14, 32, 10, 12, 14, 14, 14, 14, 14, 14, 14])
    ws.freeze_panes = f"A{FILA_HEADER + 1}"


# ══════════════════════════════════════════════════════════
# MOVIMIENTOS
# ══════════════════════════════════════════════════════════

def _hoja_movimientos(ws, movs):
    _titulo_hoja(ws, 11,
                 f"Movimientos — "
                 f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    FILA_HEADER = 3
    _encabezado(ws, FILA_HEADER, [
        "Tipo", "Código", "Producto", "Cantidad",
        "Precio Unit.", "Rebaja U.", "Total Venta",
        "Motivo", "Local", "Fecha", "Total Merma",
    ])

    por_dia = {}
    for m in movs:
        por_dia.setdefault(_solo_fecha(m["fecha"]), []).append(m)

    total_merma = total_rebajas = 0.0
    fila = FILA_HEADER

    for dia in sorted(por_dia.keys()):
        rebajas_dia = venta_dia = merma_dia = 0.0
        for m in por_dia[dia]:
            fila += 1
            cant = float(m["cantidad"] or 0)
            reb = float(m["rebaja"] or 0)
            pu = float(m["precio_unitario_momento"] or 0)
            es_salida = m["tipo"] == "SALIDA"
            es_baja = m["tipo"] == "BAJA"
            motivo_lower = (m["motivo"] or "").strip().lower()
            es_venta = es_salida and motivo_lower == "venta"
            es_merma_salida = es_salida and not es_venta
            total_venta = ""
            col_merma = ""
            if es_venta:
                total_venta = _r2((pu - reb) * cant)
                rebajas_dia += reb * cant
                venta_dia += total_venta
            elif es_baja or es_merma_salida:
                col_merma = _r2((pu - reb) * cant)
                merma_dia += col_merma

            valores = [
                m["tipo"], m.get("codigo") or "", m["producto"], cant,
                _r2(pu) if (es_salida or es_baja) else "",
                _r2(reb) if es_venta else "",
                total_venta, m["motivo"] or "", m["local"],
                m["fecha"], col_merma,
            ]
            es_merma_row = es_baja or es_merma_salida
            for col, val in enumerate(valores, 1):
                c = ws.cell(row=fila, column=col, value=val)
                c.border = BORDE_FINO
                if es_merma_row:
                    c.font = _FONT_MERMA
                if col in (1, 2, 4, 8, 9, 10):
                    c.alignment = _ALIGN_CENTER
                elif col in (5, 6, 7, 11):
                    c.alignment = _ALIGN_RIGHT
                    c.number_format = _NUM_2DEC
                else:
                    c.alignment = _ALIGN_LEFT

        total_merma += merma_dia
        total_rebajas += rebajas_dia
        fila += 1
        _cierre_dia(ws, fila, f"TOTAL DEL DÍA — {dia}",
                    rebajas_dia, venta_dia, ncols=11,
                    col_rebaja=6, col_venta=7,
                    col_merma=11, merma=merma_dia)
        fila += 1

    ws["A2"] = (f"Rebajas (acumulado): ${_r2(total_rebajas):,.2f}   |   "
                f"Merma (acumulado): ${_r2(total_merma):,.2f}")
    ws["A2"].font = _FONT_BOLD_MERMA
    ws.merge_cells("A2:K2")
    ws.row_dimensions[2].height = 20
    _anchos(ws, [16, 12, 30, 10, 12, 10, 14, 18, 16, 20, 14])
    ws.freeze_panes = f"A{FILA_HEADER + 1}"


# ══════════════════════════════════════════════════════════
# VENTAS POR LOCAL
# ══════════════════════════════════════════════════════════

def _hoja_ventas_local(ws, local, movs):
    if movs:
        mn = _solo_fecha(movs[0]["fecha"])
        mx = _solo_fecha(movs[-1]["fecha"])
        subtitulo = f"{mn} → {mx}"
    else:
        subtitulo = datetime.now().strftime('%Y-%m-%d %H:%M')

    _titulo_hoja(ws, 7,
                 f"Ventas — {local['nombre']} — {subtitulo}")
    FILA_HEADER = 3
    _encabezado(ws, FILA_HEADER, [
        "Código", "Producto", "Cantidad",
        "Precio Unit.", "Rebaja U.", "Total Venta", "Fecha",
    ])
    fila = FILA_HEADER

    if not movs:
        fila += 1
        c = ws.cell(row=fila, column=1,
                    value="Sin ventas registradas en este local.")
        c.font = _FONT_ITALIC_GREY
        c.alignment = _ALIGN_LEFT
        ws.merge_cells(start_row=fila, start_column=1,
                       end_row=fila, end_column=7)
        fila += 1
        _cierre_dia(ws, fila, "TOTAL DEL PERÍODO", 0.0, 0.0,
                    ncols=7, col_rebaja=5, col_venta=6)
        _anchos(ws, [12, 32, 10, 12, 10, 14, 20])
        return

    por_dia = {}
    for m in movs:
        por_dia.setdefault(_solo_fecha(m["fecha"]), []).append(m)

    reb_periodo = venta_periodo = 0.0
    for dia in sorted(por_dia.keys()):
        reb_dia = venta_dia = 0.0
        for m in por_dia[dia]:
            fila += 1
            cant = float(m["cantidad"] or 0)
            reb = float(m["rebaja"] or 0)
            pu = float(m["precio_unitario_momento"] or 0)
            total = _r2((pu - reb) * cant)
            reb_dia += reb * cant
            venta_dia += total

            valores = [m.get("codigo") or "", m["producto"], cant,
                       _r2(pu), _r2(reb), total, m["fecha"]]
            for col, val in enumerate(valores, 1):
                c = ws.cell(row=fila, column=col, value=val)
                c.border = BORDE_FINO
                if col in (1, 3, 7):
                    c.alignment = _ALIGN_CENTER
                elif col in (4, 5, 6):
                    c.alignment = _ALIGN_RIGHT
                    c.number_format = _NUM_2DEC
                else:
                    c.alignment = _ALIGN_LEFT

        reb_periodo += reb_dia
        venta_periodo += venta_dia
        fila += 1
        _cierre_dia(ws, fila, f"TOTAL DEL DÍA — {dia}",
                    reb_dia, venta_dia, ncols=7,
                    col_rebaja=5, col_venta=6)
        fila += 1

    fila += 1
    _cierre_dia(ws, fila, "TOTAL DEL PERÍODO",
                reb_periodo, venta_periodo, ncols=7,
                col_rebaja=5, col_venta=6)
    _anchos(ws, [12, 32, 10, 12, 10, 14, 20])
    ws.freeze_panes = f"A{FILA_HEADER + 1}"


# ══════════════════════════════════════════════════════════
# VENTAS GENERALES
# ══════════════════════════════════════════════════════════

def _hoja_ventas_generales(ws):
    ventas = _ventas_por_mes_historico()
    if ventas:
        claves = sorted(ventas.keys())
        anio_ini, mes_ini = claves[0]
        anio_fin, mes_fin = claves[-1]
        rango = (f"{NOMBRES_MES[mes_ini-1]} {anio_ini} → "
                 f"{NOMBRES_MES[mes_fin-1]} {anio_fin}")
    else:
        rango = datetime.now().strftime("%Y-%m-%d")

    _titulo_hoja(ws, 3, f"Ventas Generales — {rango}")
    FILA_HEADER = 3
    _encabezado(ws, FILA_HEADER, ["Mes", "Venta Mes", "Acumulado"])

    if not ventas:
        fila = FILA_HEADER + 1
        c = ws.cell(row=fila, column=1,
                    value="Sin ventas registradas.")
        c.font = _FONT_ITALIC_GREY
        c.alignment = _ALIGN_LEFT
        ws.merge_cells(start_row=fila, start_column=1,
                       end_row=fila, end_column=3)
        _anchos(ws, [26, 16, 16])
        return

    fila = FILA_HEADER
    acum = 0.0
    for (anio, mes) in sorted(ventas.keys()):
        v = ventas[(anio, mes)]
        acum += v
        fila += 1
        c = ws.cell(row=fila, column=1,
                    value=f"{NOMBRES_MES[mes-1]} {anio}")
        c.border = BORDE_FINO
        c.alignment = _ALIGN_LEFT
        c = ws.cell(row=fila, column=2, value=_r2(v))
        c.border = BORDE_FINO
        c.alignment = _ALIGN_RIGHT
        c.number_format = _NUM_2DEC
        c = ws.cell(row=fila, column=3, value=_r2(acum))
        c.border = BORDE_FINO
        c.alignment = _ALIGN_RIGHT
        c.number_format = _NUM_2DEC

    fila += 1
    c = ws.cell(row=fila, column=1, value="TOTAL DEL PERÍODO")
    c.font = _FONT_BOLD_TITULO
    c.alignment = _ALIGN_LEFT
    total = sum(ventas.values())
    for col in (2, 3):
        c = ws.cell(row=fila, column=col, value=_r2(total))
        c.font = _FONT_BOLD
        c.alignment = _ALIGN_RIGHT
        c.number_format = _NUM_2DEC
    _estilizar_fila(ws, fila, 3)
    _anchos(ws, [26, 16, 16])


def _ventas_por_mes_historico() -> dict:
    with get_conn() as conn:
        filas = conn.execute(
            "SELECT CAST(strftime('%Y', fecha) AS INTEGER) AS anio, "
            "CAST(strftime('%m', fecha) AS INTEGER) AS mes, "
            "SUM((precio_unitario_momento - rebaja) * cantidad) AS venta "
            "FROM movimientos WHERE tipo='SALIDA' "
            "AND LOWER(TRIM(COALESCE(motivo,'')))='venta' "
            "GROUP BY anio, mes ORDER BY anio, mes"
        ).fetchall()
    return {(int(r["anio"]), int(r["mes"])): float(r["venta"] or 0.0)
            for r in filas}


# ══════════════════════════════════════════════════════════
# EXCEL DE SALIDAS DEL DÍA
# ══════════════════════════════════════════════════════════

def generar_excel_salidas_hoy(local_id=None) -> bytes:
    """
    Bytes de un .xlsx con TODAS las SALIDA del día actual,
    SIN filtrar por motivo. Útil para cuadre diario.

    Si `local_id` es None o GENERAL_ID → todas las locales.
    Lanza ValueError si no hay ninguna salida hoy.
    """
    hoy = datetime.now().strftime("%Y-%m-%d")

    sql = (
        "SELECT m.fecha, m.cantidad, m.motivo, m.rebaja, "
        "       m.precio_unitario_momento, m.usuario, "
        "       p.nombre AS producto, p.codigo AS codigo, "
        "       l.nombre AS local "
        "FROM movimientos m "
        "JOIN productos p ON p.id=m.producto_id "
        "JOIN locales l ON l.id=m.local_id "
        "WHERE m.tipo='SALIDA' AND DATE(m.fecha)=?"
    )
    params = [hoy]
    if local_id is not None and local_id != GENERAL_ID:
        sql += " AND m.local_id=?"
        params.append(local_id)
    sql += " ORDER BY m.fecha ASC, m.id ASC"

    with get_conn() as conn:
        filas = conn.execute(sql, tuple(params)).fetchall()
    salidas = [dict(r) for r in filas]

    if not salidas:
        raise ValueError("No hay salidas registradas hoy")

    wb = Workbook()
    ws = wb.active
    ws.title = "Salidas del día"

    # ── Título ──
    titulo = f"Salidas del día — {hoy}"
    if local_id is not None and local_id != GENERAL_ID:
        titulo += f" — {salidas[0]['local']}"
    else:
        titulo += " — Todos los locales"
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13, color=COLOR_TITULO)
    ws.merge_cells("A1:I1")
    ws.row_dimensions[1].height = 26

    # ── Encabezados ──
    FILA_HEADER = 3
    _encabezado(ws, FILA_HEADER, [
        "Código", "Producto", "Cantidad", "Precio Unit.",
        "Rebaja U.", "Motivo", "Hora", "Usuario", "Local",
    ])

    # ── Filas ──
    total_unidades = 0.0
    total_valor = 0.0
    fila = FILA_HEADER
    for s in salidas:
        fila += 1
        cant = float(s["cantidad"] or 0)
        reb = float(s["rebaja"] or 0)
        pu = float(s["precio_unitario_momento"] or 0)
        total = (pu - reb) * cant
        total_unidades += cant
        total_valor += total

        partes = (s["fecha"] or "").split(" ")
        hora = partes[1][:5] if len(partes) > 1 else ""

        valores = [
            s.get("codigo") or "",
            s["producto"],
            cant,
            _r2(pu),
            _r2(reb),
            s["motivo"] or "",
            hora,
            s["usuario"],
            s.get("local") or "",
        ]
        for col, val in enumerate(valores, 1):
            c = ws.cell(row=fila, column=col, value=val)
            c.border = BORDE_FINO
            if col in (3, 4, 5, 7):
                c.alignment = _ALIGN_CENTER
                if col in (4, 5):
                    c.number_format = _NUM_2DEC
            else:
                c.alignment = _ALIGN_LEFT

    # ── Totales ──
    fila += 1
    c = ws.cell(row=fila, column=1, value="TOTAL DEL DÍA")
    c.font = _FONT_BOLD_TITULO
    c.alignment = _ALIGN_LEFT
    c.border = BORDE_FINO

    c = ws.cell(row=fila, column=3, value=total_unidades)
    c.font = _FONT_BOLD
    c.alignment = _ALIGN_CENTER
    c.border = BORDE_FINO
    c.number_format = _NUM_2DEC

    c = ws.cell(row=fila, column=4, value=round(total_valor, 2))
    c.font = _FONT_BOLD
    c.alignment = _ALIGN_CENTER
    c.border = BORDE_FINO
    c.number_format = _NUM_2DEC

    for col in (2, 5, 6, 7, 8, 9):
        ws.cell(row=fila, column=col).border = BORDE_FINO

    for col in range(1, 10):
        ws.cell(row=fila, column=col).fill = _FILL_TOTAL

    _anchos(ws, [12, 32, 10, 12, 10, 20, 8, 12, 16])
    ws.freeze_panes = f"A{FILA_HEADER + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════
# EXCEL DIARIO (resumen + una hoja por local)
# ══════════════════════════════════════════════════════════

def generar_excel_diario() -> bytes:
    """
    Excel del día actual con:
      - Hoja "Resumen": totales por local (entradas/salidas/bajas/venta/merma).
      - Una hoja por cada local activo con sus movimientos del día.
    """
    hoy = datetime.now().strftime("%Y-%m-%d")
    locales_activos = loc.listar_locales(solo_activos=True)

    wb = Workbook()

    ws_res = wb.active
    ws_res.title = "Resumen"
    _hoja_resumen_diario(ws_res, hoy, locales_activos)

    usados = {"resumen"}
    for local in locales_activos:
        nombre_hoja = _sanitizar(local["nombre"], usados)
        ws = wb.create_sheet(nombre_hoja)
        _hoja_local_diaria(ws, local, hoy)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hoja_resumen_diario(ws, hoy, locales):
    """Resumen por local con UNA sola query agrupada."""
    _titulo_hoja(ws, 7, f"Resumen diario — {hoy}")

    FILA_HEADER = 3
    _encabezado(ws, FILA_HEADER, [
        "Local", "Entradas", "Salidas", "Bajas",
        "Ventas (CUP)", "Merma (CUP)", "Total movs",
    ])

    if not locales:
        return

    ids = [l["id"] for l in locales]
    placeholders = ",".join("?" * len(ids))

    # Una sola query para todos los locales
    sql = f"""
        SELECT
            local_id,
            SUM(CASE WHEN tipo='ENTRADA' THEN 1 ELSE 0 END) AS n_ent,
            SUM(CASE WHEN tipo='SALIDA' THEN 1 ELSE 0 END) AS n_sal,
            SUM(CASE WHEN tipo='BAJA' THEN 1 ELSE 0 END) AS n_baj,
            COALESCE(SUM(CASE WHEN tipo='SALIDA'
                THEN (precio_unitario_momento - rebaja) * cantidad
                ELSE 0 END), 0) AS venta,
            COALESCE(SUM(CASE WHEN tipo='SALIDA'
                AND LOWER(TRIM(COALESCE(motivo,'')))='merma'
                THEN (precio_unitario_momento - rebaja) * cantidad
                ELSE 0 END), 0) AS merma
        FROM movimientos
        WHERE local_id IN ({placeholders})
          AND DATE(fecha)=?
        GROUP BY local_id
    """
    params = ids + [hoy]
    with get_conn() as conn:
        filas = conn.execute(sql, tuple(params)).fetchall()

    por_local = {r["local_id"]: dict(r) for r in filas}

    fila = FILA_HEADER
    tot_ent = tot_sal = tot_baj = 0
    tot_venta = tot_merma = 0.0

    for local in locales:
        fila += 1
        lid = local["id"]
        data = por_local.get(lid, {})
        ent = int(data.get("n_ent") or 0)
        sal = int(data.get("n_sal") or 0)
        baj = int(data.get("n_baj") or 0)
        venta = float(data.get("venta") or 0)
        merma = float(data.get("merma") or 0)

        valores = [
            local["nombre"], ent, sal, baj,
            _r2(venta), _r2(merma), ent + sal + baj,
        ]
        for col, val in enumerate(valores, 1):
            c = ws.cell(row=fila, column=col, value=val)
            c.border = BORDE_FINO
            if col == 1:
                c.alignment = _ALIGN_LEFT
            elif col in (2, 3, 4, 7):
                c.alignment = _ALIGN_CENTER
            else:
                c.alignment = _ALIGN_RIGHT
                c.number_format = _NUM_2DEC

        tot_ent += ent
        tot_sal += sal
        tot_baj += baj
        tot_venta += venta
        tot_merma += merma

    # Fila TOTAL
    fila += 1
    c = ws.cell(row=fila, column=1, value="TOTAL")
    c.font = _FONT_BOLD_TITULO
    c.alignment = _ALIGN_LEFT
    c.border = BORDE_FINO

    for col, val in [(2, tot_ent), (3, tot_sal),
                     (4, tot_baj), (7, tot_ent + tot_sal + tot_baj)]:
        c = ws.cell(row=fila, column=col, value=val)
        c.font = _FONT_BOLD
        c.alignment = _ALIGN_CENTER
        c.border = BORDE_FINO

    for col, val in [(5, tot_venta), (6, tot_merma)]:
        c = ws.cell(row=fila, column=col, value=_r2(val))
        c.font = _FONT_BOLD
        c.alignment = _ALIGN_RIGHT
        c.number_format = _NUM_2DEC
        c.border = BORDE_FINO

    for col in range(1, 8):
        ws.cell(row=fila, column=col).fill = _FILL_TOTAL

    _anchos(ws, [26, 11, 11, 10, 14, 14, 12])
    ws.freeze_panes = f"A{FILA_HEADER + 1}"


def _hoja_local_diaria(ws, local, hoy):
    _titulo_hoja(ws, 8, f"{local['nombre']} — {hoy}")

    with get_conn() as conn:
        filas = conn.execute(
            "SELECT m.fecha, m.tipo, m.cantidad, m.motivo, m.rebaja, "
            "m.precio_unitario_momento, m.usuario, "
            "p.nombre AS producto, p.codigo AS codigo "
            "FROM movimientos m JOIN productos p ON p.id=m.producto_id "
            "WHERE m.local_id=? AND DATE(m.fecha)=? "
            "AND m.tipo IN ('ENTRADA','SALIDA','BAJA') "
            "ORDER BY m.fecha ASC, m.id ASC",
            (local["id"], hoy)).fetchall()
    movs = [dict(r) for r in filas]

    FILA_HEADER = 3
    _encabezado(ws, FILA_HEADER, [
        "Hora", "Tipo", "Código", "Producto",
        "Cantidad", "Precio Unit.", "Motivo", "Usuario",
    ])

    if not movs:
        fila = FILA_HEADER + 1
        c = ws.cell(row=fila, column=1,
                    value="Sin movimientos registrados hoy.")
        c.font = _FONT_ITALIC_GREY
        c.alignment = _ALIGN_LEFT
        ws.merge_cells(start_row=fila, start_column=1,
                       end_row=fila, end_column=8)
        _anchos(ws, [10, 12, 12, 32, 10, 12, 20, 12])
        return

    fila = FILA_HEADER
    for m in movs:
        fila += 1
        cant = float(m["cantidad"] or 0)
        pu = float(m["precio_unitario_momento"] or 0)
        partes = (m["fecha"] or "").split(" ")
        hora = partes[1][:8] if len(partes) > 1 else ""

        valores = [
            hora,
            m["tipo"],
            m.get("codigo") or "",
            m["producto"],
            cant,
            _r2(pu) if m["tipo"] in ("SALIDA", "BAJA") else "",
            m["motivo"] or "",
            m["usuario"],
        ]
        for col, val in enumerate(valores, 1):
            c = ws.cell(row=fila, column=col, value=val)
            c.border = BORDE_FINO
            if col in (1, 2, 3, 5):
                c.alignment = _ALIGN_CENTER
            elif col == 6:
                c.alignment = _ALIGN_RIGHT
                if val != "":
                    c.number_format = _NUM_2DEC
            else:
                c.alignment = _ALIGN_LEFT

    _anchos(ws, [10, 12, 12, 32, 10, 12, 20, 12])
    ws.freeze_panes = f"A{FILA_HEADER + 1}"