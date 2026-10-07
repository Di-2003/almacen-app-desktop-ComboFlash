"""
full_test.py — Test EXHAUSTIVO del sistema multimoneda + fixes.

Verifica CADA función, CADA caso borde, CADA conversión y CADA
comportamiento del sistema multimoneda implementado en Fases 1-3.

Uso:
    python full_test.py
"""
import os
import sys
import shutil
from datetime import datetime
from pathlib import Path

# ============================================================
# Preparar entorno ANTES de importar db.py
# ============================================================
TEST_DIR = Path("./_test_full").absolute()
if TEST_DIR.exists():
    shutil.rmtree(TEST_DIR)
TEST_DIR.mkdir(parents=True)
os.environ["FLET_APP_STORAGE_DATA"] = str(TEST_DIR)

# ============================================================
# Imports
# ============================================================
from db import (inicializar_db, get_conn, GENERAL_ID,
                get_pref, set_pref, MONEDAS_VALIDAS,
                VERSION_ESQUEMA)
import inventario as inv
import locales as loc
import usuarios as um
from seguridad import crear_hash, verificar_password, generar_salt


# ============================================================
# Test runner
# ============================================================
class TestRunner:
    def __init__(self):
        self.pasados = 0
        self.fallados = 0
        self.errores = []

    def test(self, nombre, fn):
        try:
            fn()
            self.pasados += 1
            print(f"  \033[92m✓\033[0m {nombre}")
        except AssertionError as e:
            self.fallados += 1
            self.errores.append((nombre, str(e)))
            print(f"  \033[91m✗\033[0m {nombre}: {e}")
        except Exception as e:
            self.fallados += 1
            msg = f"{type(e).__name__}: {e}"
            self.errores.append((nombre, msg))
            print(f"  \033[91m💥\033[0m {nombre}: {msg}")

    def resumen(self):
        total = self.pasados + self.fallados
        print()
        print("=" * 66)
        if self.fallados == 0:
            print(f"  \033[92m✅ TODOS LOS TESTS PASARON "
                  f"({self.pasados}/{total})\033[0m")
        else:
            print(f"  \033[91m❌ {self.fallados}/{total} "
                  f"tests fallados\033[0m")
            print()
            for n, e in self.errores:
                print(f"    - {n}: {e}")
        print("=" * 66)
        return self.fallados == 0


# ============================================================
# Helpers
# ============================================================
def eq(a, b, msg=""):
    assert a == b, f"{msg}: esperado={b!r} obtenido={a!r}"

def ne(a, b, msg=""):
    assert a != b, f"{msg}: no debe ser {b!r}"

def gt(a, b, msg=""):
    assert a > b, f"{msg}: {a!r} debe ser > {b!r}"

def ge(a, b, msg=""):
    assert a >= b, f"{msg}: {a!r} debe ser >= {b!r}"

def lt(a, b, msg=""):
    assert a < b, f"{msg}: {a!r} debe ser < {b!r}"

def casi_eq(a, b, tol=0.01, msg=""):
    assert abs(float(a) - float(b)) < tol, \
        f"{msg}: esperado≈{b} obtenido={a}"

def raises(exc_type, fn, msg=""):
    try:
        fn()
    except exc_type:
        return
    except Exception as e:
        raise AssertionError(
            f"{msg}: esperaba {exc_type.__name__}, llegó "
            f"{type(e).__name__}: {e}")
    raise AssertionError(
        f"{msg}: no lanzó excepción (esperaba {exc_type.__name__})")


ADMIN = {"username": "admin", "rol": "admin"}
ALMACEN_ID = None
CENTRO_ID = None


# ============================================================
# Setup — INICIALIZA BD, ADMIN, ALMACÉN Y CENTRO
# ============================================================
def setup():
    global ALMACEN_ID, CENTRO_ID
    print("\n\033[94m[SETUP] Inicializando BD temporal…\033[0m")
    inicializar_db()
    h, s = crear_hash("admin1234")
    with get_conn() as conn:
        conn.execute(
            "INSERT INTO usuarios(username,password_hash,salt,rol,"
            "creado) VALUES(?,?,?,'admin',?)",
            ("admin", h, s,
             datetime.now().strftime("%Y-%m-%d %H:%M:%S")))

    # Almacén (auto-creado por la BD)
    alm = loc.obtener_almacen()
    if alm is None:
        raise RuntimeError("Almacén no auto-creado")
    ALMACEN_ID = alm["id"]

    # Centro (creado aquí para usarlo en los tests)
    CENTRO_ID = loc.abrir_tienda("Centro")

    print(f"  BD creada en: {TEST_DIR}")
    print(f"  Almacén id={ALMACEN_ID}  Centro id={CENTRO_ID}")


# ============================================================
# 1. Esquema v7
# ============================================================
def test_esquema_v7(t):
    def version():
        eq(VERSION_ESQUEMA, 7)
    t.test("Esquema: VERSION_ESQUEMA = 7", version)

    def monedas_validas():
        eq(MONEDAS_VALIDAS, ("CUP", "USD", "EUR"))
    t.test("Esquema: MONEDAS_VALIDAS = CUP/USD/EUR", monedas_validas)

    def columnas():
        with get_conn() as conn:
            cols = conn.execute("PRAGMA table_info(productos)").fetchall()
            nombres = {c["name"] for c in cols}
            for col in ("precio_costo_orig", "precio_unitario_orig",
                        "moneda_costo", "moneda_venta"):
                assert col in nombres, f"falta columna {col}"
    t.test("Esquema: columnas de moneda en productos", columnas)

    def defaults():
        eq(get_pref("tasa_usd"), "1.0")
        eq(get_pref("tasa_eur"), "1.0")
        eq(get_pref("moneda_visualizacion"), "CUP")
        eq(get_pref("motivo_default_salida"), "Venta")
    t.test("Esquema: defaults de monedas y tasas", defaults)

    def almacen_existe():
        assert ALMACEN_ID is not None
        alm = loc.obtener_local(ALMACEN_ID)
        eq(alm["nombre"], "Almacén")
        eq(alm["es_almacen"], 1)
    t.test("Esquema: Almacén auto-creado", almacen_existe)


# ============================================================
# 2. Conversión de moneda
# ============================================================
def test_conversion(t):
    inv.set_config("tasa_usd", "250")
    inv.set_config("tasa_eur", "300")

    def get_tasa_usd():
        casi_eq(inv.get_tasa_usd(), 250.0)
    t.test("Conversión: get_tasa_usd()", get_tasa_usd)

    def get_tasa_eur():
        casi_eq(inv.get_tasa_eur(), 300.0)
    t.test("Conversión: get_tasa_eur()", get_tasa_eur)

    def get_tasa_cup():
        casi_eq(inv.get_tasa("CUP"), 1.0)
    t.test("Conversión: get_tasa('CUP') = 1.0", get_tasa_cup)

    def get_tasa_generica():
        casi_eq(inv.get_tasa("USD"), 250.0)
        casi_eq(inv.get_tasa("EUR"), 300.0)
        casi_eq(inv.get_tasa("usd"), 250.0)
    t.test("Conversión: get_tasa() case-insensitive", get_tasa_generica)

    def a_cup():
        casi_eq(inv.a_cup(10, "USD"), 2500.0)
        casi_eq(inv.a_cup(10, "EUR"), 3000.0)
        casi_eq(inv.a_cup(10, "CUP"), 10.0)
    t.test("Conversión: a_cup()", a_cup)

    def de_cup():
        casi_eq(inv.de_cup(2500, "USD"), 10.0)
        casi_eq(inv.de_cup(3000, "EUR"), 10.0)
        casi_eq(inv.de_cup(100, "CUP"), 100.0)
    t.test("Conversión: de_cup()", de_cup)

    def convertir():
        casi_eq(inv.convertir(10, "USD", "EUR"), 10 * 250 / 300)
        casi_eq(inv.convertir(10, "USD", "USD"), 10.0)
        casi_eq(inv.convertir(10, "USD", "CUP"), 2500.0)
    t.test("Conversión: convertir() entre monedas", convertir)

    def tasa_invalida():
        inv.set_config("tasa_usd", "abc")
        casi_eq(inv.get_tasa_usd(), 1.0, msg="fallback a 1.0")
        inv.set_config("tasa_usd", "250")
    t.test("Conversión: tasa inválida → 1.0", tasa_invalida)


# ============================================================
# 3. Moneda de visualización
# ============================================================
def test_moneda_visualizacion(t):
    inv.set_config("tasa_usd", "250")
    inv.set_config("tasa_eur", "300")

    def default():
        inv.set_moneda_visualizacion("CUP")
        eq(inv.get_moneda_visualizacion(), "CUP")
    t.test("Mon. visualización: default CUP", default)

    def cambiar():
        inv.set_moneda_visualizacion("USD")
        eq(inv.get_moneda_visualizacion(), "USD")
        inv.set_moneda_visualizacion("EUR")
        eq(inv.get_moneda_visualizacion(), "EUR")
    t.test("Mon. visualización: cambiar", cambiar)

    def invalida():
        inv.set_moneda_visualizacion("XYZ")
        eq(inv.get_moneda_visualizacion(), "CUP")
    t.test("Mon. visualización: inválida → CUP", invalida)

    def mostrar_precio():
        inv.set_moneda_visualizacion("USD")
        s = inv.mostrar_precio(2500)
        assert "USD$" in s, f"esperaba USD$, obtuve {s}"
    t.test("Mon. visualización: mostrar_precio()", mostrar_precio)

    inv.set_moneda_visualizacion("CUP")


# ============================================================
# 4. Promedio ponderado
# ============================================================
def test_promedio(t):
    def simple():
        casi_eq(inv._promedio_ponderado(10, 1, 10, 2), 1.5)
    t.test("Promedio: 10@1 + 10@2 = 1.5", simple)

    def stock_cero():
        casi_eq(inv._promedio_ponderado(0, 0, 20, 5), 5.0)
    t.test("Promedio: stock 0 + 20@5 = 5", stock_cero)

    def cant_nueva_cero():
        casi_eq(inv._promedio_ponderado(10, 5, 0, 0), 5.0)
    t.test("Promedio: 10@5 + 0@0 = 5", cant_nueva_cero)

    def total_cero():
        casi_eq(inv._promedio_ponderado(0, 0, 0, 0), 0.0)
    t.test("Promedio: todo 0 = 0", total_cero)


# ============================================================
# 5. Entrada en CUP
# ============================================================
def test_entrada_cup(t):
    def nueva():
        info = inv.registrar_entrada(
            "Prod CUP 1", 50, ADMIN, ALMACEN_ID,
            codigo="FA0001", precio_costo=10, precio_unitario=20,
            moneda="CUP")
        eq(info["stock"], 50.0)
    t.test("Entrada CUP: crear nuevo", nueva)

    def check_valores():
        p = inv.buscar_producto_por_codigo("FA0001", ALMACEN_ID)
        casi_eq(p["precio_costo"], 10.0)
        casi_eq(p["precio_unitario"], 20.0)
        casi_eq(p["precio_costo_orig"], 10.0)
        casi_eq(p["precio_unitario_orig"], 20.0)
        eq(p["moneda_costo"], "CUP")
        eq(p["moneda_venta"], "CUP")
    t.test("Entrada CUP: moneda y valores correctos", check_valores)

    def sumar_stock():
        info = inv.registrar_entrada(
            "Prod CUP 1", 30, ADMIN, ALMACEN_ID,
            precio_costo=10, precio_unitario=20, moneda="CUP")
        eq(info["stock"], 80.0)
    t.test("Entrada CUP: sumar stock", sumar_stock)


# ============================================================
# 6. Entrada en USD
# ============================================================
def test_entrada_usd(t):
    inv.set_config("tasa_usd", "250")

    def nueva():
        info = inv.registrar_entrada(
            "Prod USD 1", 10, ADMIN, ALMACEN_ID,
            codigo="FB0001", precio_costo=1, precio_unitario=2,
            moneda="USD")
        eq(info["stock"], 10.0)
    t.test("Entrada USD: crear nuevo", nueva)

    def check_valores():
        p = inv.buscar_producto_por_codigo("FB0001", ALMACEN_ID)
        casi_eq(p["precio_costo_orig"], 1.0)
        casi_eq(p["precio_unitario_orig"], 2.0)
        casi_eq(p["precio_costo"], 250.0)
        casi_eq(p["precio_unitario"], 500.0)
        eq(p["moneda_costo"], "USD")
        eq(p["moneda_venta"], "USD")
    t.test("Entrada USD: valores orig vs CUP correctos", check_valores)

    def segunda_entrada_mismo_usd():
        inv.registrar_entrada(
            "Prod USD 1", 10, ADMIN, ALMACEN_ID,
            precio_costo=2, precio_unitario=3, moneda="USD")
        p = inv.buscar_producto_por_codigo("FB0001", ALMACEN_ID)
        casi_eq(p["precio_costo_orig"], 1.5, msg="promedio USD")
        casi_eq(p["precio_unitario_orig"], 3.0,
                msg="venta se sobrescribe")
        casi_eq(p["precio_costo"], 375.0, msg="1.5 USD × 250")
        casi_eq(p["precio_unitario"], 750.0, msg="3 USD × 250")
    t.test("Entrada USD: promedio ponderado + venta sobrescrita",
           segunda_entrada_mismo_usd)


# ============================================================
# 7. Entrada en EUR
# ============================================================
def test_entrada_eur(t):
    inv.set_config("tasa_eur", "300")

    def nueva():
        inv.registrar_entrada(
            "Prod EUR 1", 5, ADMIN, ALMACEN_ID,
            codigo="FC0001", precio_costo=10, precio_unitario=15,
            moneda="EUR")
        p = inv.buscar_producto_por_codigo("FC0001", ALMACEN_ID)
        casi_eq(p["precio_costo_orig"], 10.0)
        casi_eq(p["precio_costo"], 3000.0)
        eq(p["moneda_costo"], "EUR")
    t.test("Entrada EUR: crear con conversión", nueva)


# ============================================================
# 8. Entrada mixta
# ============================================================
def test_entrada_mixta(t):
    inv.set_config("tasa_usd", "250")
    inv.set_config("tasa_eur", "300")

    def product_usd_entrada_eur():
        p = inv.buscar_producto_por_codigo("FB0001", ALMACEN_ID)
        stock_antes = p["stock"]
        orig_antes = p["precio_costo_orig"]

        inv.registrar_entrada(
            "Prod USD 1", 5, ADMIN, ALMACEN_ID,
            precio_costo=10, moneda="EUR")

        p2 = inv.buscar_producto_por_codigo("FB0001", ALMACEN_ID)
        # 10 EUR × 300 = 3000 CUP ÷ 250 = 12 USD
        expected_orig = (stock_antes * orig_antes + 5 * 12.0) / (stock_antes + 5)
        casi_eq(p2["precio_costo_orig"], expected_orig,
                msg="promedio en USD tras entrada EUR")
    t.test("Entrada mixta: producto USD + entrada EUR",
           product_usd_entrada_eur)


# ============================================================
# 9. Editar precio con moneda
# ============================================================
def test_editar_precio(t):
    inv.set_config("tasa_usd", "250")
    inv.set_config("tasa_eur", "300")

    def nuevo_producto():
        inv.registrar_entrada(
            "Prod Editar 1", 10, ADMIN, ALMACEN_ID,
            codigo="FD0001", precio_costo=10, precio_unitario=20,
            moneda="CUP")
    t.test("Editar: preparar producto", nuevo_producto)

    def set_costo_usd():
        p = inv.buscar_producto_por_codigo("FD0001", ALMACEN_ID)
        inv.set_precio_costo(p["id"], 5, "USD", ADMIN)
        p2 = inv.buscar_producto_por_codigo("FD0001", ALMACEN_ID)
        casi_eq(p2["precio_costo_orig"], 5.0)
        casi_eq(p2["precio_costo"], 1250.0)
        eq(p2["moneda_costo"], "USD")
    t.test("Editar: P. Costo a USD", set_costo_usd)

    def set_venta_eur():
        p = inv.buscar_producto_por_codigo("FD0001", ALMACEN_ID)
        inv.set_precio_unitario(p["id"], 8, "EUR", ADMIN)
        p2 = inv.buscar_producto_por_codigo("FD0001", ALMACEN_ID)
        casi_eq(p2["precio_unitario_orig"], 8.0)
        casi_eq(p2["precio_unitario"], 2400.0)
        eq(p2["moneda_venta"], "EUR")
    t.test("Editar: P. Venta a EUR", set_venta_eur)

    def cambio_a_cup():
        p = inv.buscar_producto_por_codigo("FD0001", ALMACEN_ID)
        inv.set_precio_costo(p["id"], 2000, "CUP", ADMIN)
        p2 = inv.buscar_producto_por_codigo("FD0001", ALMACEN_ID)
        casi_eq(p2["precio_costo_orig"], 2000.0)
        casi_eq(p2["precio_costo"], 2000.0)
        eq(p2["moneda_costo"], "CUP")
    t.test("Editar: P. Costo a CUP", cambio_a_cup)

    def negativo():
        p = inv.buscar_producto_por_codigo("FD0001", ALMACEN_ID)
        raises(ValueError,
               lambda: inv.set_precio_costo(p["id"], -5, "USD", ADMIN))
    t.test("Editar: precio negativo rechazado", negativo)

    def propagar():
        inv.registrar_entrada(
            "Prod Editar 1", 5, ADMIN, CENTRO_ID,
            codigo="FD0001", precio_costo=10, precio_unitario=20,
            moneda="CUP")
        p_alm = inv.buscar_producto_por_codigo("FD0001", ALMACEN_ID)
        inv.set_precio_costo(p_alm["id"], 3, "USD", ADMIN)
        p_cent = inv.buscar_producto_por_codigo("FD0001", CENTRO_ID)
        casi_eq(p_cent["precio_costo_orig"], 3.0)
        eq(p_cent["moneda_costo"], "USD")
    t.test("Editar: propagación a otros locales", propagar)


# ============================================================
# 10. Cambiar moneda de un precio
# ============================================================
def test_cambiar_moneda_precio(t):
    inv.set_config("tasa_usd", "250")
    inv.set_config("tasa_eur", "300")

    def preparar():
        inv.registrar_entrada(
            "Prod Cambio Mon", 10, ADMIN, ALMACEN_ID,
            codigo="FE0001", precio_costo=5, precio_unitario=10,
            moneda="USD")
    t.test("Cambio moneda: preparar en USD", preparar)

    def usd_a_cup():
        p = inv.buscar_producto_por_codigo("FE0001", ALMACEN_ID)
        inv.cambiar_moneda_precio(p["id"], "costo", "CUP", ADMIN)
        p2 = inv.buscar_producto_por_codigo("FE0001", ALMACEN_ID)
        eq(p2["moneda_costo"], "CUP")
        casi_eq(p2["precio_costo_orig"], 1250.0)
        casi_eq(p2["precio_costo"], 1250.0)
    t.test("Cambio moneda: USD → CUP convierte", usd_a_cup)

    def cup_a_eur():
        p = inv.buscar_producto_por_codigo("FE0001", ALMACEN_ID)
        inv.cambiar_moneda_precio(p["id"], "costo", "EUR", ADMIN)
        p2 = inv.buscar_producto_por_codigo("FE0001", ALMACEN_ID)
        eq(p2["moneda_costo"], "EUR")
        casi_eq(p2["precio_costo_orig"], 1250 / 300, tol=0.01)
        casi_eq(p2["precio_costo"], 1250.0, tol=0.1)
    t.test("Cambio moneda: CUP → EUR convierte", cup_a_eur)

    def invalida():
        p = inv.buscar_producto_por_codigo("FE0001", ALMACEN_ID)
        raises(ValueError,
               lambda: inv.cambiar_moneda_precio(
                   p["id"], "costo", "XYZ", ADMIN))
    t.test("Cambio moneda: moneda inválida rechazada", invalida)

    def campo_invalido():
        p = inv.buscar_producto_por_codigo("FE0001", ALMACEN_ID)
        raises(ValueError,
               lambda: inv.cambiar_moneda_precio(
                   p["id"], "otro", "CUP", ADMIN))
    t.test("Cambio moneda: campo inválido rechazado", campo_invalido)


# ============================================================
# 11. Recálculo al cambiar tasa
# ============================================================
def test_recalcular(t):
    inv.set_config("tasa_usd", "1")
    inv.set_config("tasa_eur", "1")

    def preparar():
        inv.registrar_entrada(
            "Recalc USD", 10, ADMIN, ALMACEN_ID,
            codigo="FF0001", precio_costo=10, precio_unitario=20,
            moneda="USD")
        inv.registrar_entrada(
            "Recalc CUP", 10, ADMIN, ALMACEN_ID,
            codigo="FF0002", precio_costo=100, precio_unitario=200,
            moneda="CUP")
    t.test("Recalcular: preparar", preparar)

    def inicial():
        pu = inv.buscar_producto_por_codigo("FF0001", ALMACEN_ID)
        pc = inv.buscar_producto_por_codigo("FF0002", ALMACEN_ID)
        casi_eq(pu["precio_costo"], 10.0)
        casi_eq(pc["precio_costo"], 100.0)
    t.test("Recalcular: valores iniciales con tasa 1", inicial)

    def cambiar_tasa_y_recalcular():
        inv.set_config("tasa_usd", "5")
        n = inv.recalcular_todos_los_precios()
        ge(n, 1, "al menos 1 producto actualizado")
    t.test("Recalcular: cambiar tasa y llamar",
           cambiar_tasa_y_recalcular)

    def verificar_usd():
        pu = inv.buscar_producto_por_codigo("FF0001", ALMACEN_ID)
        casi_eq(pu["precio_costo"], 50.0)
        casi_eq(pu["precio_unitario"], 100.0)
    t.test("Recalcular: producto USD actualizado", verificar_usd)

    def verificar_cup_no_tocado():
        pc = inv.buscar_producto_por_codigo("FF0002", ALMACEN_ID)
        casi_eq(pc["precio_costo"], 100.0, msg="CUP no cambia")
        casi_eq(pc["precio_unitario"], 200.0)
    t.test("Recalcular: producto CUP no cambia",
           verificar_cup_no_tocado)

    def cambiar_tasa_eur():
        inv.set_config("tasa_eur", "8")
        inv.registrar_entrada(
            "Recalc EUR", 10, ADMIN, ALMACEN_ID,
            codigo="FF0003", precio_costo=10, precio_unitario=20,
            moneda="EUR")
        p = inv.buscar_producto_por_codigo("FF0003", ALMACEN_ID)
        casi_eq(p["precio_costo"], 80.0)
        inv.set_config("tasa_eur", "4")
        inv.recalcular_todos_los_precios()
        p2 = inv.buscar_producto_por_codigo("FF0003", ALMACEN_ID)
        casi_eq(p2["precio_costo"], 40.0)
    t.test("Recalcular: EUR también se actualiza", cambiar_tasa_eur)


# ============================================================
# 12. Salida
# ============================================================
def test_salida(t):
    inv.set_config("tasa_usd", "250")

    def preparar():
        inv.registrar_entrada(
            "Salida USD", 100, ADMIN, ALMACEN_ID,
            codigo="FG0001", precio_costo=1, precio_unitario=2,
            moneda="USD")
    t.test("Salida: preparar en USD", preparar)

    def salida_cup():
        info = inv.registrar_salida(
            "Salida USD", 10, ADMIN, ALMACEN_ID,
            motivo="Venta")
        eq(info["stock"], 90.0)
        with get_conn() as conn:
            mov = conn.execute(
                "SELECT * FROM movimientos WHERE tipo='SALIDA' "
                "ORDER BY id DESC LIMIT 1").fetchone()
            casi_eq(mov["precio_unitario_momento"], 500.0)
    t.test("Salida: precio momento en CUP", salida_cup)

    def stock_insuficiente():
        raises(inv.StockInsuficiente,
               lambda: inv.registrar_salida(
                   "Salida USD", 99999, ADMIN, ALMACEN_ID))
    t.test("Salida: stock insuficiente", stock_insuficiente)


# ============================================================
# 13. Traspaso conserva moneda
# ============================================================
def test_traspaso(t):
    inv.set_config("tasa_usd", "300")

    def preparar():
        inv.registrar_entrada(
            "Traspaso Mon", 50, ADMIN, ALMACEN_ID,
            codigo="FH0001", precio_costo=2, precio_unitario=4,
            moneda="USD")
    t.test("Traspaso: preparar en USD", preparar)

    def traspasar():
        inv.registrar_traspaso(
            "Traspaso Mon", 20, ADMIN, ALMACEN_ID, CENTRO_ID)
        pc = inv.buscar_producto_por_codigo("FH0001", CENTRO_ID)
        assert pc is not None, "producto no llegó a Centro"
        eq(pc["moneda_costo"], "USD")
        eq(pc["moneda_venta"], "USD")
        casi_eq(pc["precio_costo_orig"], 2.0)
        casi_eq(pc["precio_costo"], 600.0)
    t.test("Traspaso: conserva moneda y valores", traspasar)

    def cambio_tasa_afecta_destino():
        inv.set_config("tasa_usd", "500")
        inv.recalcular_todos_los_precios()
        pc = inv.buscar_producto_por_codigo("FH0001", CENTRO_ID)
        casi_eq(pc["precio_costo"], 1000.0)
    t.test("Traspaso: recálculo afecta destino",
           cambio_tasa_afecta_destino)


# ============================================================
# 14. Traspaso CUP no se toca al recalcular
# ============================================================
def test_traspaso_cup(t):
    inv.set_config("tasa_usd", "250")

    def preparar():
        inv.registrar_entrada(
            "Traspaso CUP", 30, ADMIN, ALMACEN_ID,
            codigo="FI0001", precio_costo=100, precio_unitario=200,
            moneda="CUP")
    t.test("Traspaso CUP: preparar", preparar)

    def traspasar():
        inv.registrar_traspaso(
            "Traspaso CUP", 10, ADMIN, ALMACEN_ID, CENTRO_ID)
        pc = inv.buscar_producto_por_codigo("FI0001", CENTRO_ID)
        eq(pc["moneda_costo"], "CUP")
        casi_eq(pc["precio_costo"], 100.0)
    t.test("Traspaso CUP: conserva CUP", traspasar)

    def recalc_no_afecta():
        inv.set_config("tasa_usd", "999")
        inv.recalcular_todos_los_precios()
        pc = inv.buscar_producto_por_codigo("FI0001", CENTRO_ID)
        casi_eq(pc["precio_costo"], 100.0, msg="CUP no cambia")
    t.test("Traspaso CUP: recálculo no lo toca", recalc_no_afecta)


# ============================================================
# 15. Excel
# ============================================================
def test_excel(t):
    def generar():
        import excel_sync as excel
        data = excel.generar_excel_bytes()
        gt(len(data), 1000, "Excel debe pesar > 1 KB")
        assert data[:2] == b"PK"
    t.test("Excel: generar bytes válidos", generar)

    def tiene_columnas():
        import excel_sync as excel
        import io
        from openpyxl import load_workbook
        data = excel.generar_excel_bytes()
        wb = load_workbook(io.BytesIO(data))
        ws = wb["Almacén"] if "Almacén" in wb.sheetnames else wb[wb.sheetnames[0]]
        headers = [ws.cell(row=3, column=c).value for c in range(1, 12)]
        assert "Moneda" in headers, f"falta Moneda, headers={headers}"
        assert "Costo orig." in headers
        assert "Costo CUP" in headers
        assert "Venta orig." in headers
        assert "Venta CUP" in headers
    t.test("Excel: columnas Moneda/Costo orig./CUP", tiene_columnas)


# ============================================================
# 16. Moneda de visualización
# ============================================================
def test_visualizacion_dashboard(t):
    def cup():
        inv.set_moneda_visualizacion("CUP")
        eq(inv.get_moneda_visualizacion(), "CUP")
    t.test("Dashboard: moneda CUP", cup)

    def usd():
        inv.set_moneda_visualizacion("USD")
        eq(inv.get_moneda_visualizacion(), "USD")
    t.test("Dashboard: moneda USD", usd)

    def eur():
        inv.set_moneda_visualizacion("EUR")
        eq(inv.get_moneda_visualizacion(), "EUR")
    t.test("Dashboard: moneda EUR", eur)
    inv.set_moneda_visualizacion("CUP")


# ============================================================
# 17. Fixes previos
# ============================================================
def test_fixes_previos(t):
    inv.set_config("motivo_default_salida", "Venta")
    inv.set_config("tasa_usd", "250")

    def motivo_venta_default():
        eq(inv.motivo_default_salida(), "Venta")
    t.test("FIX #1: motivo default Venta", motivo_venta_default)

    def entrada_existente():
        inv.registrar_entrada(
            "FIX 2", 10, ADMIN, ALMACEN_ID,
            codigo="FJ0001", precio_costo=5, precio_unitario=10,
            moneda="CUP")
        info = inv.registrar_entrada(
            "FIX 2", 20, ADMIN, ALMACEN_ID,
            codigo="FJ0001", precio_costo=5, precio_unitario=10,
            moneda="CUP")
        eq(info["stock"], 30.0)
    t.test("FIX #2: entrada de producto existente", entrada_existente)

    def totales_por_concepto():
        tc = inv.totales_por_concepto(ALMACEN_ID)
        gt(tc["entradas"]["n"], 0, "debe haber entradas")
    t.test("FIX #3: totales_por_concepto", totales_por_concepto)

    def eliminar_movimiento():
        inv.registrar_entrada(
            "FIX 5", 10, ADMIN, ALMACEN_ID,
            codigo="FK0001", precio_costo=1, precio_unitario=2,
            moneda="CUP")
        p = inv.buscar_producto_por_codigo("FK0001", ALMACEN_ID)
        eq(p["stock"], 10.0)
        with get_conn() as conn:
            mov = conn.execute(
                "SELECT id FROM movimientos WHERE tipo='ENTRADA' "
                "AND producto_id=? ORDER BY id DESC LIMIT 1",
                (p["id"],)).fetchone()
        inv.eliminar_movimiento(mov["id"], ADMIN)
        p2 = inv.buscar_producto_por_codigo("FK0001", ALMACEN_ID)
        eq(p2["stock"], 0.0)
    t.test("FIX #5: eliminar movimiento revierte stock",
           eliminar_movimiento)

    def buscar_global():
        p = inv.buscar_producto_global_por_codigo("FB0001")
        assert p is not None
    t.test("FIX #6: búsqueda global", buscar_global)

    def codigo_replicado():
        p1 = inv.buscar_producto_por_codigo("FB0001", ALMACEN_ID)
        assert p1 is not None
    t.test("FIX #6: código replicado", codigo_replicado)


# ============================================================
# 18. Validaciones de moneda inválida
# ============================================================
def test_validaciones_moneda(t):
    def entrada_moneda_invalida():
        inv.registrar_entrada(
            "Mon Invalida", 10, ADMIN, ALMACEN_ID,
            codigo="FL0001", precio_costo=10, precio_unitario=20,
            moneda="XYZ")
        p = inv.buscar_producto_por_codigo("FL0001", ALMACEN_ID)
        eq(p["moneda_costo"], "CUP")
    t.test("Validación: entrada con moneda inválida → CUP",
           entrada_moneda_invalida)

    def set_precio_moneda_invalida():
        p = inv.buscar_producto_por_codigo("FL0001", ALMACEN_ID)
        inv.set_precio_costo(p["id"], 100, "XYZ", ADMIN)
        p2 = inv.buscar_producto_por_codigo("FL0001", ALMACEN_ID)
        eq(p2["moneda_costo"], "CUP")
    t.test("Validación: set_precio con moneda inválida → CUP",
           set_precio_moneda_invalida)


# ============================================================
# 19. Múltiples productos en diferentes monedas
# ============================================================
def test_multiples_monedas(t):
    inv.set_config("tasa_usd", "100")
    inv.set_config("tasa_eur", "200")

    def crear():
        inv.registrar_entrada(
            "Multi CUP", 10, ADMIN, ALMACEN_ID,
            codigo="FM0001", precio_costo=50, precio_unitario=100,
            moneda="CUP")
        inv.registrar_entrada(
            "Multi USD", 10, ADMIN, ALMACEN_ID,
            codigo="FM0002", precio_costo=1, precio_unitario=2,
            moneda="USD")
        inv.registrar_entrada(
            "Multi EUR", 10, ADMIN, ALMACEN_ID,
            codigo="FM0003", precio_costo=1, precio_unitario=2,
            moneda="EUR")
    t.test("Multi moneda: crear 3 productos", crear)

    def verificar():
        p1 = inv.buscar_producto_por_codigo("FM0001", ALMACEN_ID)
        p2 = inv.buscar_producto_por_codigo("FM0002", ALMACEN_ID)
        p3 = inv.buscar_producto_por_codigo("FM0003", ALMACEN_ID)
        casi_eq(p1["precio_costo"], 50.0)
        casi_eq(p2["precio_costo"], 100.0)
        casi_eq(p3["precio_costo"], 200.0)
    t.test("Multi moneda: valores correctos", verificar)

    def cambiar_todas_las_tasas():
        inv.set_config("tasa_usd", "500")
        inv.set_config("tasa_eur", "800")
        inv.recalcular_todos_los_precios()
        p1 = inv.buscar_producto_por_codigo("FM0001", ALMACEN_ID)
        p2 = inv.buscar_producto_por_codigo("FM0002", ALMACEN_ID)
        p3 = inv.buscar_producto_por_codigo("FM0003", ALMACEN_ID)
        casi_eq(p1["precio_costo"], 50.0)
        casi_eq(p2["precio_costo"], 500.0)
        casi_eq(p3["precio_costo"], 800.0)
    t.test("Multi moneda: recálculo respeta cada moneda",
        cambiar_todas_las_tasas)


# ============================================================
# 20. Casos borde
# ============================================================
def test_casos_borde(t):
    inv.set_config("tasa_usd", "0")

    def a_cup_con_tasa_cero():
        casi_eq(inv.a_cup(10, "USD"), 0.0)
    t.test("Borde: a_cup con tasa 0", a_cup_con_tasa_cero)

    def de_cup_con_tasa_cero():
        result = inv.de_cup(100, "USD")
        ge(result, 0)
    t.test("Borde: de_cup con tasa 0 no explota",
        de_cup_con_tasa_cero)

    inv.set_config("tasa_usd", "250")

    def precio_cero():
        inv.registrar_entrada(
            "Borde Cero", 10, ADMIN, ALMACEN_ID,
            codigo="FN0001", precio_costo=0, precio_unitario=0,
            moneda="USD")
        p = inv.buscar_producto_por_codigo("FN0001", ALMACEN_ID)
        casi_eq(p["precio_costo"], 0.0)
        casi_eq(p["precio_costo_orig"], 0.0)
        eq(p["moneda_costo"], "USD")
    t.test("Borde: precio 0 USD", precio_cero)

    def cambiar_a_cup_cero():
        p = inv.buscar_producto_por_codigo("FN0001", ALMACEN_ID)
        inv.cambiar_moneda_precio(p["id"], "costo", "CUP", ADMIN)
        p2 = inv.buscar_producto_por_codigo("FN0001", ALMACEN_ID)
        casi_eq(p2["precio_costo_orig"], 0.0)
        eq(p2["moneda_costo"], "CUP")
    t.test("Borde: cambiar moneda de precio 0", cambiar_a_cup_cero)


# ============================================================
# 21. Integridad
# ============================================================
def test_integridad(t):
    def moneda_consistente():
        with get_conn() as conn:
            rows = conn.execute(
                "SELECT moneda_costo, moneda_venta FROM productos"
            ).fetchall()
            for r in rows:
                assert r["moneda_costo"] in MONEDAS_VALIDAS
                assert r["moneda_venta"] in MONEDAS_VALIDAS
    t.test("Integridad: monedas válidas en BD", moneda_consistente)

    def orig_igual_cup_si_cup():
        with get_conn() as conn:
            rows = conn.execute(
                "SELECT * FROM productos WHERE moneda_costo='CUP'"
            ).fetchall()
            for r in rows:
                casi_eq(r["precio_costo"], r["precio_costo_orig"],
                        tol=0.01,
                        msg=f"producto {r['id']} CUP inconsistente")
    t.test("Integridad: CUP orig == CUP valor", orig_igual_cup_si_cup)

    def orig_cup_si_usd():
        # FIX: sincronizar CUP con la tasa actual antes de verificar
        inv.set_config("tasa_usd", "250")
        inv.recalcular_todos_los_precios()
        with get_conn() as conn:
            rows = conn.execute(
                "SELECT * FROM productos WHERE moneda_costo='USD'"
            ).fetchall()
            for r in rows:
                expected = float(r["precio_costo_orig"]) * 250
                casi_eq(r["precio_costo"], expected, tol=0.5,
                        msg=f"producto {r['id']} USD inconsistente")
    t.test("Integridad: USD CUP = orig × tasa", orig_cup_si_usd)

    def orig_cup_si_eur():
        # Añado también verificación para EUR
        inv.set_config("tasa_eur", "300")
        inv.recalcular_todos_los_precios()
        with get_conn() as conn:
            rows = conn.execute(
                "SELECT * FROM productos WHERE moneda_costo='EUR'"
            ).fetchall()
            for r in rows:
                expected = float(r["precio_costo_orig"]) * 300
                casi_eq(r["precio_costo"], expected, tol=0.5,
                        msg=f"producto {r['id']} EUR inconsistente")
    t.test("Integridad: EUR CUP = orig × tasa", orig_cup_si_eur)


# ============================================================
# 22. Excel avanzado
# ============================================================
def test_excel_avanzado(t):
    def generar_con_multi_moneda():
        import excel_sync as excel
        data = excel.generar_excel_bytes()
        gt(len(data), 1000)
    t.test("Excel: generar con multi-moneda",
           generar_con_multi_moneda)

    def hojas_presentes():
        import excel_sync as excel
        import io
        from openpyxl import load_workbook
        data = excel.generar_excel_bytes()
        wb = load_workbook(io.BytesIO(data))
        assert "Movimientos" in wb.sheetnames
        assert "VentasGenerales" in wb.sheetnames
    t.test("Excel: hojas fijas presentes", hojas_presentes)


# ============================================================
# MAIN
# ============================================================
def main():
    print()
    print("=" * 66)
    print("  \033[93mFULL TEST — Multimoneda + Fixes\033[0m")
    print("  Verifica TODAS las funciones y casos borde.")
    print("=" * 66)

    setup()
    t = TestRunner()

    print("\n\033[94m[1] Esquema v7\033[0m")
    test_esquema_v7(t)
    print("\n\033[94m[2] Conversión de moneda\033[0m")
    test_conversion(t)
    print("\n\033[94m[3] Moneda de visualización\033[0m")
    test_moneda_visualizacion(t)
    print("\n\033[94m[4] Promedio ponderado\033[0m")
    test_promedio(t)
    print("\n\033[94m[5] Entrada en CUP\033[0m")
    test_entrada_cup(t)
    print("\n\033[94m[6] Entrada en USD\033[0m")
    test_entrada_usd(t)
    print("\n\033[94m[7] Entrada en EUR\033[0m")
    test_entrada_eur(t)
    print("\n\033[94m[8] Entrada mixta\033[0m")
    test_entrada_mixta(t)
    print("\n\033[94m[9] Editar precio con moneda\033[0m")
    test_editar_precio(t)
    print("\n\033[94m[10] Cambiar moneda de un precio\033[0m")
    test_cambiar_moneda_precio(t)
    print("\n\033[94m[11] Recálculo al cambiar tasa\033[0m")
    test_recalcular(t)
    print("\n\033[94m[12] Salida\033[0m")
    test_salida(t)
    print("\n\033[94m[13] Traspaso conserva moneda\033[0m")
    test_traspaso(t)
    print("\n\033[94m[14] Traspaso CUP no se toca\033[0m")
    test_traspaso_cup(t)
    print("\n\033[94m[15] Excel con columnas de moneda\033[0m")
    test_excel(t)
    print("\n\033[94m[16] Moneda de visualización\033[0m")
    test_visualizacion_dashboard(t)
    print("\n\033[94m[17] Fixes previos\033[0m")
    test_fixes_previos(t)
    print("\n\033[94m[18] Validaciones de moneda inválida\033[0m")
    test_validaciones_moneda(t)
    print("\n\033[94m[19] Múltiples monedas simultáneas\033[0m")
    test_multiples_monedas(t)
    print("\n\033[94m[20] Casos borde\033[0m")
    test_casos_borde(t)
    print("\n\033[94m[21] Integridad de datos\033[0m")
    test_integridad(t)
    print("\n\033[94m[22] Excel avanzado\033[0m")
    test_excel_avanzado(t)

    print(f"\n\033[94m[CLEANUP] Borrando {TEST_DIR}…\033[0m")
    try:
        shutil.rmtree(TEST_DIR, ignore_errors=True)
        print("  OK")
    except Exception as e:
        print(f"  Aviso: {e}")

    ok = t.resumen()
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()